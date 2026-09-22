import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Master Execution Guide: WESAD ECG Stress Pipeline")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_after = Pt(2)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("End-to-End Operational Protocol: From 700 Hz Wearable Biosignal Telemetry to ML Benchmarks & Web Deployment")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    p_sub.paragraph_format.space_after = Pt(6)
    
    # Metadata
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("Author: Mukesh Yadav (ECE, JSSATEN)  |  Dataset: WESAD (15 Subjects)  |  Validation: 15-Fold LOSO  |  ROC-AUC: 0.9494")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    p_meta.paragraph_format.space_after = Pt(14)
    
    def add_sec_heading(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(3, 105, 161)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        return h

    def add_step_box(step_num, title, badge_text, desc, cmd, output_text=None):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f"Step {step_num}: {title} ")
        r1.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(15, 23, 42)
        
        r_badge = p.add_run(f"[{badge_text}]")
        r_badge.bold = True
        r_badge.font.size = Pt(8.5)
        r_badge.font.color.rgb = RGBColor(2, 132, 199)
        
        p_desc = cell.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(3)
        r_desc = p_desc.add_run(desc)
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = RGBColor(51, 65, 85)
        
        p_cmd = cell.add_paragraph()
        p_cmd.paragraph_format.space_after = Pt(3)
        r_cmd = p_cmd.add_run(cmd)
        r_cmd.font.name = "Consolas"
        r_cmd.font.size = Pt(9)
        r_cmd.font.bold = True
        r_cmd.font.color.rgb = RGBColor(15, 23, 42)
        
        if output_text:
            p_out = cell.add_paragraph()
            p_out.paragraph_format.space_after = Pt(0)
            r_out_lbl = p_out.add_run("Output: ")
            r_out_lbl.bold = True
            r_out_lbl.font.size = Pt(8.5)
            r_out = p_out.add_run(output_text)
            r_out.font.size = Pt(8.5)
            r_out.font.italic = True
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)

    # Section 1
    add_sec_heading("1. System Architecture & Information Flow")
    p_arch = doc.add_paragraph()
    r_arch = p_arch.add_run("""Raw 700 Hz Lead-II Chest ECG (WESAD, 15 Subjects)
  └──> Step 0: Python Extraction to S*_ECG_labels.mat
        ├──> MATLAB DSP & Feature Pipeline:
        │     • 0.5–40 Hz 4th-order zero-phase Butterworth filtering (filtfilt)
        │     • Pan-Tompkins & MAD-adaptive R-peak detection (refractory lockout = 350 ms)
        │     • 60-second sliding windows with 50% overlap (30s step) -> 445 windows
        │     • 8 Core HRV Metrics + Subject-Specific Relative Calibration (Δx)
        │     • Master 6-Panel Research Dashboard & 300 DPI Medical Figures
        └──> Python ML Benchmark & Deployment:
              • 15-Fold Leave-One-Subject-Out (LOSO) Cross-Validation
              • 6 Machine Learning Architectures (LR, SVM, RF, ET, HistGB, MLP)
              • Explainability (Permutation Importance & Clinical Odds Ratios)
              • Interactive Streamlit Clinical Telemetry Web Application""")
    r_arch.font.name = "Consolas"
    r_arch.font.size = Pt(8.5)
    r_arch.font.color.rgb = RGBColor(30, 41, 59)
    p_arch.paragraph_format.space_after = Pt(10)

    # Section 2
    add_sec_heading("2. Prerequisites & Environment Setup")
    p_env = doc.add_paragraph()
    p_env.add_run("• Python 3.10+ with packages: streamlit, plotly, pandas, numpy, scipy, scikit-learn.\n"
                  "• MATLAB R2022b+ with Signal Processing Toolbox and Statistics and Machine Learning Toolbox.\n"
                  "• PowerShell command to configure environment:\n")
    p_env.paragraph_format.space_after = Pt(2)
    p_cmd_env = doc.add_paragraph()
    r_ce = p_cmd_env.add_run("cd \"C:\\Users\\YASH\\Desktop\\projects\\RESEARCH PROJECTS\\ECG_STRESS_DETECTION\"\n"
                            "python -m venv venv; .\\venv\\Scripts\\Activate.ps1; pip install -r requirements.txt")
    r_ce.font.name = "Consolas"
    r_ce.font.size = Pt(9)
    r_ce.font.bold = True
    p_cmd_env.paragraph_format.space_after = Pt(10)

    # Section 3
    add_sec_heading("3. Step-by-Step Chronological Execution Sequence")
    
    add_step_box("0", "Raw Data Ingestion & Formatting", "PYTHON",
                 "Ingests WESAD pickle files (S2.pkl - S17.pkl), extracts 700 Hz chest ECG and affective ground-truth labels, downcasts precision to float32/uint8, and saves compressed MATLAB arrays.",
                 "python python/extract_ecg_labels.py",
                 "data/processed/S*_ECG_labels.mat for all 15 subjects")
                 
    add_step_box("1", "MATLAB Workspace & Path Initialization", "MATLAB",
                 "Registers all project subdirectories, DSP functions, and feature extraction routines to the MATLAB search path.",
                 "cd 'C:\\Users\\YASH\\Desktop\\projects\\RESEARCH PROJECTS\\ECG_STRESS_DETECTION\\matlab'; setup_project",
                 "All subfolders (01_data_inspection to 06_final_results) added to path")
                 
    add_step_box("2", "Multi-Subject DSP & Feature Extraction", "MATLAB",
                 "Applies zero-phase 4th-order Butterworth bandpass (0.5-40 Hz), computes MAD dynamic noise-floor prominence, enforces 350 ms lockout, gates intervals to 300-1500 ms (40-200 BPM), and extracts 8 core HRV metrics normalized against resting baseline.",
                 "TEN_process_all_subjects",
                 "results/WESAD_HRV_features_expanded.csv (445 standardized windows)")
                 
    add_step_box("3", "Statistical Hypothesis Testing & Audits", "MATLAB",
                 "Runs paired Student's t-tests and Wilcoxon signed-rank tests confirming significant sympathetic activation and parasympathetic withdrawal (p < 0.001).",
                 "FOURTEEN_statistical_tests; FIFTEEN_feature_analysis",
                 "Statistical significance metrics and feature distribution tables")
                 
    add_step_box("4", "Publication Figures & Master Clinical Dashboard", "MATLAB",
                 "Generates 300 DPI vector figures (confusion matrix, ROC curve, timeline) and the comprehensive 6-panel clinical research board.",
                 "TWENTY_SEVEN_final_results; TWENTY_EIGHT_report_figures; TWENTY_NINE_project_dashboard",
                 "results/figures/FINAL_Project_Dashboard.png and FIG1-FIG5")
                 
    add_step_box("5", "Interactive Telemetry MATLAB Demo GUI", "MATLAB GUI",
                 "Interactive multi-tabbed clinical telemetry viewer for Lead-II ECG, full timeline, Pan-Tompkins QRS detection, and biological feature contrasts.",
                 "DEMO_stress_detection",
                 "Interactive GUI window with 4 inspection tabs")
                 
    add_step_box("6", "15-Fold LOSO Machine Learning Benchmark", "PYTHON",
                 "Performs strict 15-fold Leave-One-Subject-Out cross-validation across 6 machine learning architectures with zero subject data leakage (StandardScaler fit strictly on training fold).",
                 "python python/train_loso_ml_benchmark.py",
                 "Full performance table (ROC-AUC > 0.937 across all 6 models, MLP ROC-AUC = 0.938)")
                 
    add_step_box("7", "Explainability & Feature Importance", "PYTHON",
                 "Computes permutation importance and clinical standardized odds ratios, confirming that cardiac interval compression (ΔMeanRR) and heart rate acceleration (ΔMeanHR) govern the decision boundary.",
                 "python python/explainability_feature_importance.py",
                 "Permutation importance rankings and clinical odds ratios")
                 
    add_step_box("8", "Multi-Model Evaluation Plots & Curves", "PYTHON",
                 "Renders high-resolution comparative ROC curves, Precision-Recall curves, confusion matrices, and performance bar charts across all benchmarked models.",
                 "python python/plot_ml_evaluation.py",
                 "results/figures/python/ml_benchmark_roc_curves.png, etc.")
                 
    add_step_box("9", "Interactive Clinical Telemetry Web Application", "STREAMLIT WEB APP",
                 "Launches the browser-based clinical telemetry dashboard with real-time waveform inspection, dynamic threshold slider, and live stress needle gauge.",
                 "streamlit run demo/app.py",
                 "Local web dashboard at http://localhost:8501")

    # Section 4: Quick Run One-Liners
    add_sec_heading("4. Quick-Run Cheat Sheet")
    t_quick = doc.add_table(rows=4, cols=2)
    t_quick.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Workflow Target", "Command"]
    for i, h in enumerate(headers):
        cell = t_quick.cell(0, i)
        cell.text = h
        set_cell_background(cell, "E2E8F0")
        cell.paragraphs[0].runs[0].bold = True
        set_cell_margins(cell, 80, 80, 120, 120)
        
    rows_data = [
        ("Run Full Python ML Suite", "python python/train_loso_ml_benchmark.py; python python/explainability_feature_importance.py; python python/plot_ml_evaluation.py"),
        ("Run Headless MATLAB Pipeline", "matlab -batch \"cd('matlab'); setup_project; TWENTY_SEVEN_final_results; TWENTY_EIGHT_report_figures; TWENTY_NINE_project_dashboard; exit\""),
        ("Launch Web Dashboard", "streamlit run demo/app.py")
    ]
    for r_idx, (tgt, cmd) in enumerate(rows_data, start=1):
        c1 = t_quick.cell(r_idx, 0)
        c2 = t_quick.cell(r_idx, 1)
        c1.text = tgt
        c2.text = cmd
        c1.paragraphs[0].runs[0].font.size = Pt(8.5)
        c1.paragraphs[0].runs[0].bold = True
        c2.paragraphs[0].runs[0].font.name = "Consolas"
        c2.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_margins(c2, 60, 60, 100, 100)
        if r_idx % 2 == 1:
            set_cell_background(c1, "F8FAFC")
            set_cell_background(c2, "F8FAFC")

    # Section 5: Verification Checklist
    add_sec_heading("5. Verification Checklist & Expected Benchmarks")
    t_ver = doc.add_table(rows=6, cols=3)
    t_ver.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_v = ["Target Output", "Verification Metric", "Status"]
    for i, h in enumerate(headers_v):
        cell = t_ver.cell(0, i)
        cell.text = h
        set_cell_background(cell, "E2E8F0")
        cell.paragraphs[0].runs[0].bold = True
        set_cell_margins(cell, 80, 80, 120, 120)
        
    ver_data = [
        ("S*_ECG_labels.mat", "15 subject files, 700 Hz Lead-II ECG, synchronized ground truth labels", "Verified (Done)"),
        ("WESAD_HRV_features_expanded.csv", "445 total windows (285 baseline, 160 acute TSST stress), 8 core HRV features", "Verified (Done)"),
        ("Primary Calibrated Classifier", "ROC-AUC = 0.9494, Sensitivity = 86.25%, Specificity = 95.79% (at tau = 0.35)", "Verified (Done)"),
        ("Multi-Model ML Benchmark", "All 6 architectures (LR, SVM, RF, ET, HistGB, MLP) achieve ROC-AUC > 0.937 on held-out subjects", "Verified (Done)"),
        ("Master Dashboard Graphic", "results/figures/FINAL_Project_Dashboard.png generated at 300 DPI", "Verified (Done)")
    ]
    for r_idx, (tgt, metric, stat) in enumerate(ver_data, start=1):
        c1 = t_ver.cell(r_idx, 0)
        c2 = t_ver.cell(r_idx, 1)
        c3 = t_ver.cell(r_idx, 2)
        c1.text = tgt
        c2.text = metric
        c3.text = stat
        c1.paragraphs[0].runs[0].font.size = Pt(8.5)
        c1.paragraphs[0].runs[0].bold = True
        c2.paragraphs[0].runs[0].font.size = Pt(8.5)
        c3.paragraphs[0].runs[0].font.size = Pt(8.5)
        c3.paragraphs[0].runs[0].bold = True
        c3.paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 101, 52)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_margins(c2, 60, 60, 100, 100)
        set_cell_margins(c3, 60, 60, 100, 100)
        if r_idx % 2 == 1:
            set_cell_background(c1, "F8FAFC")
            set_cell_background(c2, "F8FAFC")
            set_cell_background(c3, "F8FAFC")

    # Output file
    out_path = r"C:\Users\YASH\Desktop\projects\RESEARCH PROJECTS\ECG_STRESS_DETECTION\EXECUTION_GUIDE.docx"
    doc.save(out_path)
    print(f"SUCCESS: Saved Word docx to {out_path}")

if __name__ == "__main__":
    create_docx()
